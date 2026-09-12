import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def generate_tasks_excel(tasks, completed_task_ids, file_path="data/tasks.xlsx"):
    """
    Generate a beautifully formatted Excel file for tasks.
    tasks: list of tuples (task_id, category, title, description, target_time)
    completed_task_ids: list of completed task_ids
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Nhiệm Vụ Hàng Ngày"
    
    # Define styles
    header_font = Font(name='Segoe UI', bold=True, color='FFFFFF', size=12)
    header_fill = PatternFill(start_color='4F81BD', end_color='4F81BD', fill_type='solid')
    header_alignment = Alignment(horizontal='center', vertical='center')
    
    cell_font = Font(name='Segoe UI', size=11)
    cell_alignment_center = Alignment(horizontal='center', vertical='center', wrap_text=True)
    cell_alignment_left = Alignment(horizontal='left', vertical='center', wrap_text=True)
    
    thin_border = Border(
        left=Side(style='thin', color='BFBFBF'),
        right=Side(style='thin', color='BFBFBF'),
        top=Side(style='thin', color='BFBFBF'),
        bottom=Side(style='thin', color='BFBFBF')
    )
    
    completed_fill = PatternFill(start_color='E2EFDA', end_color='E2EFDA', fill_type='solid')
    pending_fill = PatternFill(start_color='FCE4D6', end_color='FCE4D6', fill_type='solid')
    
    # Headers
    headers = ["ID", "Trạng Thái", "Deadline", "Thể Loại", "Tên Nhiệm Vụ", "Mô Tả Chi Tiết"]
    ws.append(headers)
    
    # Apply header styles
    for col_num, cell in enumerate(ws[1], 1):
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = thin_border
    
    # Add data
    for row_num, t in enumerate(tasks, 2):
        task_id, category, title, description, target_time = t
        
        status = "✅ Hoàn Thành" if task_id in completed_task_ids else "⏳ Đang Chờ"
        row_data = [task_id, status, target_time or "N/A", category, title, description]
        
        ws.append(row_data)
        
        # Apply styles to data rows
        for col_num, cell in enumerate(ws[row_num], 1):
            cell.font = cell_font
            cell.border = thin_border
            
            # Alignments
            if col_num in [1, 2, 3]:  # ID, Status, Deadline
                cell.alignment = cell_alignment_center
            else:  # Category, Title, Description
                cell.alignment = cell_alignment_left
                
            # Row coloring based on status
            if task_id in completed_task_ids:
                cell.fill = completed_fill
            else:
                cell.fill = pending_fill
    
    # Adjust column widths
    column_widths = {
        'A': 8,   # ID
        'B': 18,  # Trạng Thái
        'C': 12,  # Deadline
        'D': 20,  # Thể Loại
        'E': 35,  # Tên Nhiệm Vụ
        'F': 55   # Mô Tả Chi Tiết
    }
    
    for col_letter, width in column_widths.items():
        ws.column_dimensions[col_letter].width = width
        
    # Adjust row heights for word wrap
    for row in ws.iter_rows(min_row=2, max_col=6):
        ws.row_dimensions[row[0].row].height = 40
    
    # Save directory
    os.makedirs(os.path.dirname(file_path), exist_ok=True)
    wb.save(file_path)
    return file_path
